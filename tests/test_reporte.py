import pandas as pd
import pytest
from openpyxl import load_workbook

from reporte import cargar, exportar, resumir

CSV = """fecha,sucursal,vendedor,producto,cantidad,precio_unitario
2026-01-05,CDMX,Ana,Taza,2,150
2026-01-20,Monterrey,Luis,Café 1kg,1,320
2026-02-03,CDMX,Ana,Café 1kg,3,320
"""


@pytest.fixture
def ventas(tmp_path):
    ruta = tmp_path / "ventas.csv"
    ruta.write_text(CSV, encoding="utf-8")
    return cargar(str(ruta))


def test_totales_por_mes_y_vendedor(ventas):
    tablas = resumir(ventas)

    assert tablas["Por mes"].to_dict("records") == [
        {"mes": "2026-01", "ventas": 620, "unidades": 3},
        {"mes": "2026-02", "ventas": 960, "unidades": 3},
    ]
    assert tablas["Por vendedor"].iloc[0].to_dict() == {"vendedor": "Ana", "ventas": 1260, "unidades": 5}
    assert tablas["Mes x sucursal"].set_index("mes").loc["2026-01", "Monterrey"] == 320


def test_csv_sin_columnas_da_error_claro(tmp_path):
    ruta = tmp_path / "malo.csv"
    ruta.write_text("fecha,producto\n2026-01-01,Taza\n", encoding="utf-8")
    with pytest.raises(ValueError, match="cantidad"):
        cargar(str(ruta))


def test_excel_tiene_hojas_y_grafica(ventas, tmp_path):
    salida = tmp_path / "reporte.xlsx"
    exportar(resumir(ventas), str(salida))

    libro = load_workbook(salida)
    assert libro.sheetnames == ["Por mes", "Por producto", "Por vendedor", "Mes x sucursal", "Detalle"]
    assert len(libro["Por mes"]._charts) == 1
