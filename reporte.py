"""Convierte un CSV de ventas en un reporte de Excel con resúmenes y gráfica.

Uso: python reporte.py ventas.csv --salida reporte_ventas.xlsx
"""
import argparse

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNAS_REQUERIDAS = {"fecha", "sucursal", "vendedor", "producto", "cantidad", "precio_unitario"}
FORMATO_PESOS = '"$"#,##0.00'
COLOR_ENCABEZADO = "1F4E78"


def cargar(ruta: str) -> pd.DataFrame:
    df = pd.read_csv(ruta, parse_dates=["fecha"])
    faltantes = COLUMNAS_REQUERIDAS - set(df.columns)
    if faltantes:
        raise ValueError(f"Al CSV le faltan columnas: {', '.join(sorted(faltantes))}")
    return df.assign(total=df["cantidad"] * df["precio_unitario"], mes=df["fecha"].dt.to_period("M").astype(str))


def resumir(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """Arma las tablas del reporte, una por hoja."""
    def por(columna: str) -> pd.DataFrame:
        return (df.groupby(columna, as_index=False)
                  .agg(ventas=("total", "sum"), unidades=("cantidad", "sum"))
                  .sort_values("ventas", ascending=False, ignore_index=True))

    return {
        "Por mes": por("mes").sort_values("mes", ignore_index=True),
        "Por producto": por("producto"),
        "Por vendedor": por("vendedor"),
        "Mes x sucursal": df.pivot_table(index="mes", columns="sucursal", values="total",
                                         aggfunc="sum", fill_value=0).reset_index(),
        "Detalle": df.drop(columns="mes"),
    }


def dar_formato(hoja) -> None:
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
    for col in hoja.iter_cols(min_row=2):
        if hoja.cell(1, col[0].column).value in {"ventas", "total", "precio_unitario"} or hoja.title == "Mes x sucursal":
            for celda in col:
                if isinstance(celda.value, (int, float)):
                    celda.number_format = FORMATO_PESOS
    for i, col in enumerate(hoja.columns, start=1):
        ancho = max(len(str(c.value or "")) for c in col) + 2
        hoja.column_dimensions[get_column_letter(i)].width = min(ancho, 40)
    hoja.freeze_panes = "A2"


def agregar_grafica(hoja) -> None:
    grafica = BarChart()
    grafica.title = "Ventas por mes"
    grafica.y_axis.title = "Ventas ($)"
    datos = Reference(hoja, min_col=2, min_row=1, max_row=hoja.max_row)
    meses = Reference(hoja, min_col=1, min_row=2, max_row=hoja.max_row)
    grafica.add_data(datos, titles_from_data=True)
    grafica.set_categories(meses)
    hoja.add_chart(grafica, "E2")


def exportar(tablas: dict[str, pd.DataFrame], salida: str) -> None:
    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        for nombre, tabla in tablas.items():
            tabla.to_excel(writer, sheet_name=nombre, index=False)
            dar_formato(writer.sheets[nombre])
        agregar_grafica(writer.sheets["Por mes"])


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("csv")
    parser.add_argument("--salida", default="reporte_ventas.xlsx")
    args = parser.parse_args()

    tablas = resumir(cargar(args.csv))
    exportar(tablas, args.salida)
    print(f"Listo: {args.salida} ({len(tablas)} hojas)")


if __name__ == "__main__":
    main()
